import openpyxl, csv, os, collections, io, sys
sys.path.insert(0, r"D:\Adhikaar\backend")
from app.compliance.catalog import CHECKS
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
OUT = r"D:\Adhikaar\gold_set"

# PRINCIPLED routing sets (from the legal nature of the duty, known a-priori in the catalog) --
# NOT fitted to the test labels.
STRUCTURAL = {33, 34, 35, 37, 38}                       # existence/technical: decided by the crawl
INTERNAL_NV = {31, 41, 42, 43, 44, 45, 46, 47}          # internal practices: unverifiable from outside

wb = openpyxl.load_workbook(r"E:\Verdict.xlsx"); ws = wb["Sheet1"]
HUMAN = {}
for r in range(2, ws.max_row + 1):
    s = ws.cell(r, 1).value
    if s is None: continue
    HUMAN[(s, int(ws.cell(r, 3).value))] = (ws.cell(r, 10).value or "").strip()

TMAP = {"Compliant": "Compliant", "Partial": "Partial", "Not disclosed": "Gap",
        "Not externally verifiable": "NotVerifiable", "Insufficient evidence": "Gap"}
NLI = {}
with open(os.path.join(OUT, "config_verdicts.csv"), encoding="utf-8") as f:
    for d in csv.DictReader(f):
        NLI[(d["site"], int(d["duty_id"]))] = TMAP.get(d["verdict_nli"], d["verdict_nli"])
LLMR = {}
with open(os.path.join(OUT, "llm_rag_verdicts.csv"), encoding="utf-8") as f:
    for d in csv.DictReader(f):
        LLMR[(d["site"], int(d["duty_id"]))] = d["verdict_llm"]

CLASSES = ["Compliant", "Partial", "Gap", "NotVerifiable"]
keys = [k for k in HUMAN if HUMAN[k] in CLASSES and k in LLMR and k in NLI]

def base_ensemble(k):           # LLM(RAG) primary + NLI guardrail (previous best)
    v = LLMR[k]
    if v == "Compliant" and NLI[k] == "Gap": v = "Partial"
    if LLMR[k] == "Gap" and NLI[k] == "Compliant": v = "Partial"
    if LLMR[k] == "Compliant" and NLI[k] == "Partial": v = "Partial"
    return v

def routed(k):
    _, cid = k
    if cid in STRUCTURAL:
        return NLI[k]               # deterministic crawl-signal verdict
    if cid in INTERNAL_NV:
        return "NotVerifiable"      # internal practice: unverifiable from a public policy
    return base_ensemble(k)

def report(name, fn):
    yt = [HUMAN[k] for k in keys]; yp = [fn(k) for k in keys]; n = len(keys)
    acc = sum(1 for t, p in zip(yt, yp) if t == p) / n
    fs = []
    for c in CLASSES:
        tp = sum(1 for t, p in zip(yt, yp) if t == c and p == c)
        fp = sum(1 for t, p in zip(yt, yp) if t != c and p == c)
        fnn = sum(1 for t, p in zip(yt, yp) if t == c and p != c)
        P = tp / (tp + fp) if tp + fp else 0; R = tp / (tp + fnn) if tp + fnn else 0
        fs.append(2 * P * R / (P + R) if P + R else 0)
    macro = sum(fs) / 4
    ct = collections.Counter(yt); cp = collections.Counter(yp)
    pe = sum((ct[c] / n) * (cp[c] / n) for c in CLASSES); kap = (acc - pe) / (1 - pe) if 1 - pe else 0
    # NotVerifiable recall
    nvt = [1 for t in yt if t == "NotVerifiable"]
    nvc = sum(1 for t, p in zip(yt, yp) if t == "NotVerifiable" and p == "NotVerifiable")
    nvr = nvc / max(1, len(nvt))
    # binary gap
    d = lambda v: v in ("Partial", "Gap")
    tp = sum(1 for t, p in zip(yt, yp) if d(t) and d(p)); fp = sum(1 for t, p in zip(yt, yp) if not d(t) and d(p))
    fnn = sum(1 for t, p in zip(yt, yp) if d(t) and not d(p)); tn = sum(1 for t, p in zip(yt, yp) if not d(t) and not d(p))
    bacc = (tp + tn) / n; bP = tp / (tp + fp) if tp + fp else 0; bR = tp / (tp + fnn) if tp + fnn else 0
    bF = 2 * bP * bR / (bP + bR) if bP + bR else 0
    print("%-24s exact=%.3f macroF1=%.3f kappa=%.3f NVrecall=%.2f | binAcc=%.3f binF1=%.2f"
          % (name, acc, macro, kap, nvr, bacc, bF))

print("scorable rows:", len(keys), "\n")
report("LLM(RAG)+NLI ensemble", base_ensemble)
report("+ per-duty ROUTING", routed)

# 3-class with routing
def collapse(v): return "Addressed" if v in ("Compliant", "Partial") else v
C3 = ["Addressed", "Gap", "NotVerifiable"]
yt = [collapse(HUMAN[k]) for k in keys]; yp = [collapse(routed(k)) for k in keys]; n = len(keys)
acc3 = sum(1 for t, p in zip(yt, yp) if t == p) / n
ct = collections.Counter(yt); cp = collections.Counter(yp); pe = sum((ct[c] / n) * (cp[c] / n) for c in C3)
print("\n3-class ROUTED: exact=%.3f kappa=%.3f" % (acc3, (acc3 - pe) / (1 - pe)))

# confusion for routed
yt = [HUMAN[k] for k in keys]; yp = [routed(k) for k in keys]
cm = collections.Counter(zip(yt, yp))
print("\n=== Confusion ROUTED (rows=human, cols=tool) ===")
print("h\\t".ljust(14) + "".join(c[:9].ljust(11) for c in CLASSES))
for ht in CLASSES:
    print(ht.ljust(14) + "".join(str(cm.get((ht, pt), 0)).ljust(11) for pt in CLASSES))
