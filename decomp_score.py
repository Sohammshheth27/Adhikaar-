import openpyxl, csv, os, collections, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
OUT = r"D:\Adhikaar\gold_set"

wb = openpyxl.load_workbook(r"E:\Verdict.xlsx"); ws = wb["Sheet1"]
HUMAN = {}
for r in range(2, ws.max_row + 1):
    s = ws.cell(r, 1).value
    if s is None: continue
    HUMAN[(s, int(ws.cell(r, 3).value))] = (ws.cell(r, 10).value or "").strip()

DEC = {}
with open(os.path.join(OUT, "decomp_verdicts.csv"), encoding="utf-8") as f:
    for d in csv.DictReader(f):
        DEC[(d["site"], int(d["duty_id"]))] = d["verdict"]
TMAP = {"Compliant": "Compliant", "Partial": "Partial", "Not disclosed": "Gap",
        "Not externally verifiable": "NotVerifiable", "Insufficient evidence": "Gap"}
NLI = {}
with open(os.path.join(OUT, "config_verdicts.csv"), encoding="utf-8") as f:
    for d in csv.DictReader(f):
        NLI[(d["site"], int(d["duty_id"]))] = TMAP.get(d["verdict_nli"], d["verdict_nli"])

CLASSES = ["Compliant", "Partial", "Gap", "NotVerifiable"]
keys = [k for k in HUMAN if HUMAN[k] in CLASSES and k in DEC]

def dec_nli(k):   # decomposition + NLI guardrail on over-credit
    v = DEC[k]
    if v == "Compliant" and NLI[k] == "Gap": v = "Partial"
    return v

def report(name, fn):
    yt = [HUMAN[k] for k in keys]; yp = [fn(k) for k in keys]; n = len(keys)
    acc = sum(1 for t, p in zip(yt, yp) if t == p) / n
    fs = []
    for c in CLASSES:
        tp = sum(1 for t, p in zip(yt, yp) if t == c and p == c)
        fp = sum(1 for t, p in zip(yt, yp) if t != c and p == c)
        fnn = sum(1 for t, p in zip(yt, yp) if t == c and p != c)
        P = tp / (tp + fp) if tp + fp else 0; R = tp / (tp + fnn) if tp + fnn else 0
        fs.append((c, round(P, 3), round(R, 3), round(2 * P * R / (P + R) if P + R else 0, 3), tp + fnn))
    macro = sum(x[3] for x in fs) / 4
    ct = collections.Counter(yt); cp = collections.Counter(yp)
    pe = sum((ct[c] / n) * (cp[c] / n) for c in CLASSES); kap = (acc - pe) / (1 - pe) if 1 - pe else 0
    nvt = sum(1 for t in yt if t == "NotVerifiable"); nvc = sum(1 for t, p in zip(yt, yp) if t == "NotVerifiable" == p)
    d = lambda v: v in ("Partial", "Gap")
    tp = sum(1 for t, p in zip(yt, yp) if d(t) and d(p)); fp = sum(1 for t, p in zip(yt, yp) if not d(t) and d(p))
    fnn = sum(1 for t, p in zip(yt, yp) if d(t) and not d(p)); tn = sum(1 for t, p in zip(yt, yp) if not d(t) and not d(p))
    bacc = (tp + tn) / n; bP = tp / (tp + fp) if tp + fp else 0; bR = tp / (tp + fnn) if tp + fnn else 0
    bF = 2 * bP * bR / (bP + bR) if bP + bR else 0; fpr = fp / (fp + tn) if fp + tn else 0
    print("%-26s exact=%.3f macroF1=%.3f kappa=%.3f | binAcc=%.3f binF1=%.2f FPR=%.2f"
          % (name, acc, macro, kap, bacc, bF, fpr))
    return fs

print("scorable rows:", len(keys), "\n")
report("decomposition alone", lambda k: DEC[k])
fs = report("decomposition + NLI guard", dec_nli)

def collapse(v): return "Addressed" if v in ("Compliant", "Partial") else v
C3 = ["Addressed", "Gap", "NotVerifiable"]
yt = [collapse(HUMAN[k]) for k in keys]; yp = [collapse(dec_nli(k)) for k in keys]; n = len(keys)
a3 = sum(1 for t, p in zip(yt, yp) if t == p) / n
ct = collections.Counter(yt); cp = collections.Counter(yp); pe = sum((ct[c] / n) * (cp[c] / n) for c in C3)
print("\n3-class (decomp+NLI): exact=%.3f kappa=%.3f" % (a3, (a3 - pe) / (1 - pe)))

print("\n=== per-class (decomp + NLI guard) ===")
for c, P, R, F, sup in fs:
    print("  %-14s P=%.3f R=%.3f F1=%.3f support=%d" % (c, P, R, F, sup))
yt = [HUMAN[k] for k in keys]; yp = [dec_nli(k) for k in keys]
cm = collections.Counter(zip(yt, yp))
print("\n=== Confusion (rows=human, cols=tool) ===")
print("h\\t".ljust(14) + "".join(c[:9].ljust(11) for c in CLASSES))
for ht in CLASSES:
    print(ht.ljust(14) + "".join(str(cm.get((ht, pt), 0)).ljust(11) for pt in CLASSES))
