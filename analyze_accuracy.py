import openpyxl, csv, os, collections, math, json, io, sys

OUT = r"D:\Adhikaar\gold_set"
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

# ---- human gold (map to canonical vocab; drop NotApplicable from scoring) ----
wb = openpyxl.load_workbook(r"E:\Verdict.xlsx")
ws = wb["Sheet1"]
HUMAN = {}
for r in range(2, ws.max_row + 1):
    site = ws.cell(r, 1).value
    if site is None:
        continue
    HUMAN[(site, int(ws.cell(r, 3).value))] = (ws.cell(r, 10).value or "").strip()

# ---- tool verdicts (map tool vocab -> human vocab) ----
TOOLMAP = {"Compliant": "Compliant", "Partial": "Partial",
           "Not disclosed": "Gap", "Not externally verifiable": "NotVerifiable",
           "Insufficient evidence": "Gap"}
TOOL = {}
with open(os.path.join(OUT, "config_verdicts.csv"), encoding="utf-8") as f:
    for d in csv.DictReader(f):
        k = (d["site"], int(d["duty_id"]))
        TOOL[k] = {cfg: TOOLMAP.get(d["verdict_" + cfg], d["verdict_" + cfg])
                   for cfg in ("rules", "semantic", "nli")}

CLASSES = ["Compliant", "Partial", "Gap", "NotVerifiable"]

def prf(y_true, y_pred, classes):
    out = {}
    for c in classes:
        tp = sum(1 for t, p in zip(y_true, y_pred) if t == c and p == c)
        fp = sum(1 for t, p in zip(y_true, y_pred) if t != c and p == c)
        fn = sum(1 for t, p in zip(y_true, y_pred) if t == c and p != c)
        P = tp / (tp + fp) if tp + fp else 0.0
        R = tp / (tp + fn) if tp + fn else 0.0
        F = 2 * P * R / (P + R) if P + R else 0.0
        out[c] = (round(P, 3), round(R, 3), round(F, 3), tp + fn)  # support = true count
    macro = round(sum(out[c][2] for c in classes) / len(classes), 3)
    return out, macro

def kappa(y_true, y_pred, classes):
    n = len(y_true)
    po = sum(1 for t, p in zip(y_true, y_pred) if t == p) / n
    ct = collections.Counter(y_true); cp = collections.Counter(y_pred)
    pe = sum((ct[c] / n) * (cp[c] / n) for c in classes)
    return round((po - pe) / (1 - pe), 3) if (1 - pe) else 0.0

# ---- scorable rows: human verdict is one of the 4 tool-expressible classes ----
keys = [k for k in HUMAN if HUMAN[k] in CLASSES and k in TOOL]
n_na = sum(1 for k in HUMAN if HUMAN[k] == "NotApplicable")
print("scorable rows:", len(keys), "| NotApplicable excluded:", n_na, "| total:", len(HUMAN))

report = {"scorable": len(keys), "not_applicable_excluded": n_na}
for cfg in ("rules", "semantic", "nli"):
    yt = [HUMAN[k] for k in keys]
    yp = [TOOL[k][cfg] for k in keys]
    acc = sum(1 for t, p in zip(yt, yp) if t == p) / len(keys)
    perclass, macro = prf(yt, yp, CLASSES)
    kap = kappa(yt, yp, CLASSES)
    # binary gap-detection: positive = "deficient" = {Partial, Gap}; negative = {Compliant, NotVerifiable}
    def defic(v): return v in ("Partial", "Gap")
    tp = sum(1 for t, p in zip(yt, yp) if defic(t) and defic(p))
    fp = sum(1 for t, p in zip(yt, yp) if not defic(t) and defic(p))
    fn = sum(1 for t, p in zip(yt, yp) if defic(t) and not defic(p))
    tn = sum(1 for t, p in zip(yt, yp) if not defic(t) and not defic(p))
    bP = tp / (tp + fp) if tp + fp else 0.0
    bR = tp / (tp + fn) if tp + fn else 0.0
    bF = 2 * bP * bR / (bP + bR) if bP + bR else 0.0
    fpr = fp / (fp + tn) if fp + tn else 0.0
    report[cfg] = {
        "exact_accuracy": round(acc, 3), "macro_f1": macro, "cohen_kappa_tool_vs_human": kap,
        "per_class_PRF_support": {c: perclass[c] for c in CLASSES},
        "binary_gap_detection": {"precision": round(bP, 3), "recall": round(bR, 3),
                                 "f1": round(bF, 3), "false_positive_rate": round(fpr, 3),
                                 "tp": tp, "fp": fp, "fn": fn, "tn": tn},
    }

# confusion matrix for the deployed NLI config
cm = collections.Counter((HUMAN[k], TOOL[k]["nli"]) for k in keys)
print("\n=== Confusion matrix (rows=human, cols=tool NLI) ===")
hdr = "human\\tool".ljust(14) + "".join(c[:9].ljust(11) for c in CLASSES)
print(hdr)
for ht in CLASSES:
    line = ht.ljust(14) + "".join(str(cm.get((ht, pt), 0)).ljust(11) for pt in CLASSES)
    print(line)

print("\n=== SUMMARY ===")
for cfg in ("rules", "semantic", "nli"):
    rc = report[cfg]
    print(f"[{cfg}] exact-acc={rc['exact_accuracy']} macro-F1={rc['macro_f1']} "
          f"kappa={rc['cohen_kappa_tool_vs_human']} | gap-detect "
          f"P={rc['binary_gap_detection']['precision']} R={rc['binary_gap_detection']['recall']} "
          f"F1={rc['binary_gap_detection']['f1']} FPR={rc['binary_gap_detection']['false_positive_rate']}")
print("\n=== Per-class P/R/F1 (deployed NLI) ===")
for c in CLASSES:
    P, R, F, sup = report["nli"]["per_class_PRF_support"][c]
    print(f"  {c.ljust(14)} P={P} R={R} F1={F} support={sup}")

with open(os.path.join(OUT, "accuracy_results.json"), "w", encoding="utf-8") as f:
    json.dump(report, f, indent=2)
print("\nwrote accuracy_results.json")
