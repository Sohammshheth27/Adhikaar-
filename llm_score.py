import openpyxl, csv, os, collections, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
OUT = r"D:\Adhikaar\gold_set"

# human gold
wb = openpyxl.load_workbook(r"E:\Verdict.xlsx"); ws = wb["Sheet1"]
HUMAN = {}
for r in range(2, ws.max_row + 1):
    s = ws.cell(r, 1).value
    if s is None: continue
    HUMAN[(s, int(ws.cell(r, 3).value))] = (ws.cell(r, 10).value or "").strip()

# tool configs
TOOLMAP = {"Compliant": "Compliant", "Partial": "Partial", "Not disclosed": "Gap",
           "Not externally verifiable": "NotVerifiable", "Insufficient evidence": "Gap"}
TOOL = collections.defaultdict(dict)
with open(os.path.join(OUT, "config_verdicts.csv"), encoding="utf-8") as f:
    for d in csv.DictReader(f):
        k = (d["site"], int(d["duty_id"]))
        for cfg in ("rules", "semantic", "nli"):
            TOOL[k][cfg] = TOOLMAP.get(d["verdict_" + cfg], d["verdict_" + cfg])
with open(os.path.join(OUT, "llm_verdicts.csv"), encoding="utf-8") as f:
    for d in csv.DictReader(f):
        TOOL[(d["site"], int(d["duty_id"]))]["llm"] = d["verdict_llm"]

CLASSES = ["Compliant", "Partial", "Gap", "NotVerifiable"]
keys = [k for k in HUMAN if HUMAN[k] in CLASSES and "llm" in TOOL.get(k, {})]

def prf(yt, yp):
    out = {}
    for c in CLASSES:
        tp = sum(1 for t, p in zip(yt, yp) if t == c and p == c)
        fp = sum(1 for t, p in zip(yt, yp) if t != c and p == c)
        fn = sum(1 for t, p in zip(yt, yp) if t == c and p != c)
        P = tp / (tp + fp) if tp + fp else 0.0
        R = tp / (tp + fn) if tp + fn else 0.0
        F = 2 * P * R / (P + R) if P + R else 0.0
        out[c] = (round(P, 3), round(R, 3), round(F, 3), tp + fn)
    macro = round(sum(out[c][2] for c in CLASSES) / 4, 3)
    return out, macro

def kappa(yt, yp):
    n = len(yt); po = sum(1 for t, p in zip(yt, yp) if t == p) / n
    ct = collections.Counter(yt); cp = collections.Counter(yp)
    pe = sum((ct[c] / n) * (cp[c] / n) for c in CLASSES)
    return round((po - pe) / (1 - pe), 3) if 1 - pe else 0.0

def metrics(cfg):
    yt = [HUMAN[k] for k in keys]; yp = [TOOL[k][cfg] for k in keys]
    acc = sum(1 for t, p in zip(yt, yp) if t == p) / len(keys)
    per, macro = prf(yt, yp); kap = kappa(yt, yp)
    defic = lambda v: v in ("Partial", "Gap")
    tp = sum(1 for t, p in zip(yt, yp) if defic(t) and defic(p))
    fp = sum(1 for t, p in zip(yt, yp) if not defic(t) and defic(p))
    fn = sum(1 for t, p in zip(yt, yp) if defic(t) and not defic(p))
    tn = sum(1 for t, p in zip(yt, yp) if not defic(t) and not defic(p))
    bP = tp / (tp + fp) if tp + fp else 0.0
    bR = tp / (tp + fn) if tp + fn else 0.0
    bF = 2 * bP * bR / (bP + bR) if bP + bR else 0.0
    fpr = fp / (fp + tn) if fp + tn else 0.0
    return acc, macro, kap, per, (round(bP, 3), round(bR, 3), round(bF, 3), round(fpr, 3))

print("scorable rows:", len(keys))
print("\n%-10s %8s %9s %7s | %s" % ("config", "exact", "macroF1", "kappa", "gap-detect P/R/F1/FPR"))
for cfg in ("rules", "semantic", "nli", "llm"):
    acc, macro, kap, per, b = metrics(cfg)
    print("%-10s %8.3f %9.3f %7.3f | P=%.2f R=%.2f F1=%.2f FPR=%.2f"
          % (cfg, acc, macro, kap, b[0], b[1], b[2], b[3]))

print("\n=== Per-class P/R/F1 (LLM Llama-3.3-70B) ===")
_, _, _, per, _ = metrics("llm")
for c in CLASSES:
    P, R, F, sup = per[c]
    print("  %-14s P=%.3f R=%.3f F1=%.3f support=%d" % (c, P, R, F, sup))

yt = [HUMAN[k] for k in keys]; yp = [TOOL[k]["llm"] for k in keys]
cm = collections.Counter(zip(yt, yp))
print("\n=== Confusion (rows=human, cols=LLM) ===")
print("human\\llm".ljust(14) + "".join(c[:9].ljust(11) for c in CLASSES))
for ht in CLASSES:
    print(ht.ljust(14) + "".join(str(cm.get((ht, pt), 0)).ljust(11) for pt in CLASSES))
