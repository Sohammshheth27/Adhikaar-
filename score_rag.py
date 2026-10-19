import openpyxl, csv, os, collections, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
OUT = r"D:\Adhikaar\gold_set"

wb = openpyxl.load_workbook(r"E:\Verdict.xlsx"); ws = wb["Sheet1"]
HUMAN = {}
for r in range(2, ws.max_row + 1):
    s = ws.cell(r, 1).value
    if s is None: continue
    HUMAN[(s, int(ws.cell(r, 3).value))] = (ws.cell(r, 10).value or "").strip()

TMAP = {"Compliant": "Compliant", "Partial": "Partial", "Not disclosed": "Gap",
        "Not externally verifiable": "NotVerifiable", "Insufficient evidence": "Gap"}
NLI = {}
with open(os.path.join(OUT, "config_verdicts.csv"), encoding="utf-8") as f:
    for d in csv.DictReader(f):
        NLI[(d["site"], int(d["duty_id"]))] = TMAP.get(d["verdict_nli"], d["verdict_nli"])
def load_llm(fn):
    M = {}
    with open(os.path.join(OUT, fn), encoding="utf-8") as f:
        for d in csv.DictReader(f):
            M[(d["site"], int(d["duty_id"]))] = d["verdict_llm"]
    return M
LLM = load_llm("llm_verdicts.csv")
LLMR = load_llm("llm_rag_verdicts.csv")

CLASSES = ["Compliant", "Partial", "Gap", "NotVerifiable"]
keys = [k for k in HUMAN if HUMAN[k] in CLASSES and k in LLMR and k in NLI and k in LLM]

def metrics(getp):
    yt = [HUMAN[k] for k in keys]; yp = [getp(k) for k in keys]; n = len(keys)
    acc = sum(1 for t, p in zip(yt, yp) if t == p) / n
    fs = []
    for c in CLASSES:
        tp = sum(1 for t, p in zip(yt, yp) if t == c and p == c)
        fp = sum(1 for t, p in zip(yt, yp) if t != c and p == c)
        fn = sum(1 for t, p in zip(yt, yp) if t == c and p != c)
        P = tp / (tp + fp) if tp + fp else 0; R = tp / (tp + fn) if tp + fn else 0
        fs.append(2 * P * R / (P + R) if P + R else 0)
    macro = sum(fs) / 4
    ct = collections.Counter(yt); cp = collections.Counter(yp)
    pe = sum((ct[c] / n) * (cp[c] / n) for c in CLASSES); kap = (acc - pe) / (1 - pe) if 1 - pe else 0
    d = lambda v: v in ("Partial", "Gap")
    tp = sum(1 for t, p in zip(yt, yp) if d(t) and d(p)); fp = sum(1 for t, p in zip(yt, yp) if not d(t) and d(p))
    fnn = sum(1 for t, p in zip(yt, yp) if d(t) and not d(p)); tn = sum(1 for t, p in zip(yt, yp) if not d(t) and not d(p))
    bP = tp / (tp + fp) if tp + fp else 0; bR = tp / (tp + fnn) if tp + fnn else 0
    bF = 2 * bP * bR / (bP + bR) if bP + bR else 0; fpr = fp / (fp + tn) if fp + tn else 0
    bacc = (tp + tn) / n
    return acc, macro, kap, bacc, bP, bR, bF, fpr

def ens(getp, k):  # getp primary + NLI guardrails (v3)
    v = getp(k)
    if v == "Compliant" and NLI[k] == "Gap": v = "Partial"
    if getp(k) == "Gap" and NLI[k] == "Compliant": v = "Partial"
    if getp(k) == "Compliant" and NLI[k] == "Partial": v = "Partial"
    return v

print("scorable rows:", len(keys), "\n")
print("%-28s %6s %7s %6s | %6s %5s %5s %5s %5s" % ("config", "exact", "macroF1", "kappa", "binAcc", "P", "R", "F1", "FPR"))
def line(name, getp):
    a = metrics(getp)
    print("%-28s %6.3f %7.3f %6.3f | %6.3f %5.2f %5.2f %5.2f %5.2f" % (name, a[0], a[1], a[2], a[3], a[4], a[5], a[6], a[7]))
line("LLM (truncated)", lambda k: LLM[k])
line("LLM + RAG retrieval", lambda k: LLMR[k])
line("LLM(RAG) + NLI ensemble", lambda k: ens(lambda kk: LLMR[kk], k))

# 3-class
def collapse(v): return "Addressed" if v in ("Compliant", "Partial") else v
C3 = ["Addressed", "Gap", "NotVerifiable"]
def m3(getp):
    yt = [collapse(HUMAN[k]) for k in keys]; yp = [collapse(getp(k)) for k in keys]; n = len(keys)
    acc = sum(1 for t, p in zip(yt, yp) if t == p) / n
    ct = collections.Counter(yt); cp = collections.Counter(yp); pe = sum((ct[c] / n) * (cp[c] / n) for c in C3)
    return acc, (acc - pe) / (1 - pe) if 1 - pe else 0
print("\n--- 3-class (Addressed / Gap / NotVerifiable) ---")
for name, gp in [("LLM(RAG)", lambda k: LLMR[k]), ("LLM(RAG)+NLI ens", lambda k: ens(lambda kk: LLMR[kk], k))]:
    a, kp = m3(gp); print("%-20s exact=%.3f kappa=%.3f" % (name, a, kp))

# per-class for best
print("\n=== per-class, LLM(RAG)+NLI ensemble ===")
yt = [HUMAN[k] for k in keys]; yp = [ens(lambda kk: LLMR[kk], k) for k in keys]
for c in CLASSES:
    tp = sum(1 for t, p in zip(yt, yp) if t == c and p == c)
    fp = sum(1 for t, p in zip(yt, yp) if t != c and p == c)
    fn = sum(1 for t, p in zip(yt, yp) if t == c and p != c)
    P = tp / (tp + fp) if tp + fp else 0; R = tp / (tp + fn) if tp + fn else 0
    F = 2 * P * R / (P + R) if P + R else 0
    print("  %-14s P=%.3f R=%.3f F1=%.3f support=%d" % (c, P, R, F, tp + fn))
