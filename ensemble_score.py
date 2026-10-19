import openpyxl, csv, os, collections, io, sys
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
OUT = r"D:\Adhikaar\gold_set"

wb = openpyxl.load_workbook(r"E:\Verdict.xlsx"); ws = wb["Sheet1"]
HUMAN = {}
for r in range(2, ws.max_row + 1):
    s = ws.cell(r, 1).value
    if s is None: continue
    HUMAN[(s, int(ws.cell(r, 3).value))] = (ws.cell(r, 10).value or "").strip()

TMAP = {"Compliant": "Compliant", "Partial": "Partial", "Not disclosed": "Gap",
        "Not externally verifiable": "NotVerifiable", "Insufficient evidence": "Gap"}
NLI = {}
with open(os.path.join(OUT, "config_verdicts.csv"), encoding="utf-8") as f:
    for d in csv.DictReader(f):
        NLI[(d["site"], int(d["duty_id"]))] = TMAP.get(d["verdict_nli"], d["verdict_nli"])
LLM = {}
with open(os.path.join(OUT, "llm_verdicts.csv"), encoding="utf-8") as f:
    for d in csv.DictReader(f):
        LLM[(d["site"], int(d["duty_id"]))] = d["verdict_llm"]

CLASSES = ["Compliant", "Partial", "Gap", "NotVerifiable"]
keys = [k for k in HUMAN if HUMAN[k] in CLASSES and k in LLM and k in NLI]

def ens_v1(k):   # LLM primary + anti-over-credit: Compliant w/o NLI support -> Partial
    v = LLM[k]
    if v == "Compliant" and NLI[k] in ("Gap",):
        return "Partial"
    return v

def ens_v2(k):   # v1 + anti-miss: LLM Gap but NLI Compliant -> Partial
    v = ens_v1(k)
    if LLM[k] == "Gap" and NLI[k] == "Compliant":
        return "Partial"
    return v

def ens_v3(k):   # v2 + NLI confirms Partial where LLM said Compliant
    v = ens_v2(k)
    if LLM[k] == "Compliant" and NLI[k] == "Partial":
        return "Partial"
    return v

def score(name, fn):
    yt = [HUMAN[k] for k in keys]; yp = [fn(k) for k in keys]
    n = len(keys)
    acc = sum(1 for t, p in zip(yt, yp) if t == p) / n
    # macro F1
    fs = []
    for c in CLASSES:
        tp = sum(1 for t, p in zip(yt, yp) if t == c and p == c)
        fp = sum(1 for t, p in zip(yt, yp) if t != c and p == c)
        fn_ = sum(1 for t, p in zip(yt, yp) if t == c and p != c)
        P = tp / (tp + fp) if tp + fp else 0; R = tp / (tp + fn_) if tp + fn_ else 0
        fs.append(2 * P * R / (P + R) if P + R else 0)
    macro = sum(fs) / 4
    ct = collections.Counter(yt); cp = collections.Counter(yp)
    po = acc; pe = sum((ct[c] / n) * (cp[c] / n) for c in CLASSES)
    kap = (po - pe) / (1 - pe) if 1 - pe else 0
    # binary gap-detection
    d = lambda v: v in ("Partial", "Gap")
    tp = sum(1 for t, p in zip(yt, yp) if d(t) and d(p)); fp = sum(1 for t, p in zip(yt, yp) if not d(t) and d(p))
    fnn = sum(1 for t, p in zip(yt, yp) if d(t) and not d(p)); tn = sum(1 for t, p in zip(yt, yp) if not d(t) and not d(p))
    bP = tp / (tp + fp) if tp + fp else 0; bR = tp / (tp + fnn) if tp + fnn else 0
    bF = 2 * bP * bR / (bP + bR) if bP + bR else 0; fpr = fp / (fp + tn) if fp + tn else 0
    bacc = (tp + tn) / n
    print("%-26s exact=%.3f macroF1=%.3f kappa=%.3f | binary acc=%.3f P=%.2f R=%.2f F1=%.2f FPR=%.2f"
          % (name, acc, macro, kap, bacc, bP, bR, bF, fpr))

print("scorable rows:", len(keys), "\n")
score("LLM alone", lambda k: LLM[k])
score("NLI alone", lambda k: NLI[k])
score("Ensemble v1 (anti-overcredit)", ens_v1)
score("Ensemble v2 (+anti-miss)", ens_v2)
score("Ensemble v3 (+NLI Partial)", ens_v3)

# 3-class collapse (merge Compliant+Partial = 'Addressed'): removes the subjective boundary
print("\n--- 3-class (Addressed = Compliant+Partial / Gap / NotVerifiable) ---")
def collapse(v): return "Addressed" if v in ("Compliant", "Partial") else v
C3 = ["Addressed", "Gap", "NotVerifiable"]
def score3(name, fn):
    yt = [collapse(HUMAN[k]) for k in keys]; yp = [collapse(fn(k)) for k in keys]
    n = len(keys); acc = sum(1 for t, p in zip(yt, yp) if t == p) / n
    ct = collections.Counter(yt); cp = collections.Counter(yp)
    pe = sum((ct[c] / n) * (cp[c] / n) for c in C3); kap = (acc - pe) / (1 - pe) if 1 - pe else 0
    print("%-26s exact=%.3f kappa=%.3f" % (name, acc, kap))
score3("LLM alone", lambda k: LLM[k])
score3("Ensemble v2", ens_v2)
