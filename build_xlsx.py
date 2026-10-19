import os, re, csv, glob, textwrap
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\Adhikaar\gold_set"
XLSX = os.path.join(OUT, "ADHIKAAR_grading.xlsx")

HOST_MAP = {
    "jnu_ac_in": "jnu.ac.in", "du_ac_in": "du.ac.in", "iima_ac_in": "iima.ac.in",
    "akshayapatra_org": "akshayapatra.org", "tatatrusts_org": "tatatrusts.org",
    "savethechildren_in": "savethechildren.in", "robu_in": "robu.in",
    "networkintelligence_ai": "networkintelligence.ai", "powersoft_in": "powersoft.in",
}
# excel sheet-name safe + <=31 chars; map host -> sheet name
def sheet_name(host):
    s = re.sub(r"[\[\]\*\?/\\:]", "_", host)
    return s[:31]

wb = openpyxl.Workbook()

HDR = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="2F5496")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
THIN = Border(*[Side(style="thin", color="DDDDDD")] * 4)

# ---------- Sheet 1: START HERE ----------
ws0 = wb.active
ws0.title = "START HERE"
lines = [
    ("ADHIKAAR grading workbook", 16, True),
    ("", 11, False),
    ("What to do: on the 'Labels' sheet, fill the VERDICT column for each row.", 12, True),
    ("Each row = one DPDP duty applied to one website. There are 423 rows (9 sites x 47 duties).", 11, False),
    ("", 11, False),
    ("How to grade one row:", 12, True),
    ("1. Look at the 'site' in the row. Open that site's tab at the bottom (e.g. the 'robu.in' tab) to read its privacy policy.", 11, False),
    ("2. Read the 'requirement' in the row (what the law asks for).", 11, False),
    ("3. Click the VERDICT cell. A dropdown appears. Pick ONE:", 11, False),
    ("        Compliant      = the policy clearly does this", 11, False),
    ("        Partial        = mentioned, but vague or incomplete", 11, False),
    ("        Gap            = required, but missing from the policy", 11, False),
    ("        NotApplicable  = this rule does not apply to this site", 11, False),
    ("        NotVerifiable  = cannot tell from the public policy alone", 11, False),
    ("4. (Optional) Type the sentence you based it on into 'evidence_note'.", 11, False),
    ("", 11, False),
    ("Judge ONLY from the policy text and the law. Do not use any tool output. That keeps your grading independent.", 11, True),
    ("", 11, False),
    ("Tip: you can grade just one site first (47 rows). Use the filter arrow on the 'site' column in the Labels sheet.", 11, False),
    ("Tip: rows alternate shading by site so you can see where one site ends and the next begins.", 11, False),
    ("", 11, False),
    ("When done: save the file and send it back. If a second person grades too, they use a fresh copy of this same file.", 11, True),
]
r = 1
for text, size, bold in lines:
    c = ws0.cell(row=r, column=1, value=text)
    c.font = Font(size=size, bold=bold)
    r += 1
ws0.column_dimensions["A"].width = 110

# ---------- per-site policy sheets ----------
site_sheet = {}
for path in sorted(glob.glob(os.path.join(OUT, "policy_*.txt"))):
    m = re.match(r"policy_(edu|ngo|sme)_(.+)\.txt$", os.path.basename(path))
    if not m:
        continue
    host = HOST_MAP.get(m.group(2))
    if not host:
        continue
    with open(path, encoding="utf-8") as f:
        content = f.read()
    body = content.split("\n\n", 1)[1] if "\n\n" in content else content
    sn = sheet_name(host)
    site_sheet[host] = sn
    ws = wb.create_sheet(title=sn)
    ws.cell(row=1, column=1, value="PRIVACY POLICY TEXT: " + host).font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value="(read this to grade the '%s' rows on the Labels sheet)" % host).font = Font(italic=True, color="666666")
    # write policy as wrapped chunks, one paragraph-ish block per row
    rr = 4
    for para in re.split(r"\n{1,}", body):
        para = para.strip()
        if not para:
            continue
        for chunk in textwrap.wrap(para, 2000) or [para]:
            cc = ws.cell(row=rr, column=1, value=chunk[:32000])
            cc.alignment = WRAP
            rr += 1
    ws.column_dimensions["A"].width = 120

# ---------- Labels sheet ----------
wsL = wb.create_sheet(title="Labels")
cols = ["site", "sector", "duty_id", "requirement", "anchor", "in_force", "verdict", "evidence_note", "read_policy"]
for j, name in enumerate(cols, start=1):
    c = wsL.cell(row=1, column=j, value=name)
    c.font = HDR
    c.fill = HDR_FILL
    c.alignment = TOP

with open(os.path.join(OUT, "gold_labels_blank.csv"), encoding="utf-8") as f:
    data = list(csv.DictReader(f))

# shading per site
palette = ["FFFFFF", "EEF3FB"]
site_order = []
for d in data:
    if d["site"] not in site_order:
        site_order.append(d["site"])
site_shade = {s: palette[i % 2] for i, s in enumerate(site_order)}

rr = 2
for d in data:
    vals = [d["site"], d["sector"], int(d["duty_id"]), d["requirement"], d["anchor"],
            d["in_force"], "", "", None]
    for j, v in enumerate(vals, start=1):
        c = wsL.cell(row=rr, column=j, value=v)
        c.alignment = WRAP if j in (4,) else TOP
        c.border = THIN
        c.fill = PatternFill("solid", fgColor=site_shade[d["site"]])
    # hyperlink to the policy sheet
    sn = site_sheet.get(d["site"])
    if sn:
        hc = wsL.cell(row=rr, column=9, value="read %s policy" % d["site"])
        hc.value = '=HYPERLINK("#\'%s\'!A1","read policy")' % sn
        hc.font = Font(color="0563C1", underline="single")
        hc.fill = PatternFill("solid", fgColor=site_shade[d["site"]])
        hc.border = THIN
    rr += 1

# column widths
widths = {"A": 22, "B": 8, "C": 7, "D": 58, "E": 16, "F": 8, "G": 16, "H": 40, "I": 16}
for col, w in widths.items():
    wsL.column_dimensions[col].width = w
wsL.freeze_panes = "A2"
wsL.auto_filter.ref = "A1:I%d" % (rr - 1)

# dropdown validation on verdict column (G)
dv = DataValidation(type="list",
                    formula1='"Compliant,Partial,Gap,NotApplicable,NotVerifiable"',
                    allow_blank=True, showDropDown=False)
dv.prompt = "Pick one verdict"
dv.promptTitle = "Verdict"
wsL.add_data_validation(dv)
dv.add("G2:G%d" % (rr - 1))

# make Labels the first visible sheet after START HERE
wb.move_sheet("Labels", offset=-(len(wb.sheetnames) - 2))
wb.active = wb.sheetnames.index("START HERE")

wb.save(XLSX)
print("WROTE", XLSX, "rows:", rr - 2, "sheets:", wb.sheetnames)
